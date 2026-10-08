from fastapi.responses import FileResponse
from fastapi import Depends, FastAPI, HTTPException, Request
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import Response
from sqlalchemy import (
    Column,
    Integer,
    String,
    Text,
    create_engine
)
from sqlalchemy.orm import (
    Session,
    sessionmaker
)
from sqlalchemy.orm import declarative_base
from sqlalchemy import inspect, text
from sqlalchemy import LargeBinary
from pydantic import BaseModel
from typing import Literal, Optional
from pathlib import Path
import secrets
from datetime import date, datetime
from fastapi import Header
from urllib.parse import quote, unquote
from io import BytesIO, StringIO
import csv
import re
from zipfile import BadZipFile


# ============================================================
# FASTAPI APPLICATION
# ============================================================

app = FastAPI(
    title="Report Generation and Approval System",
    description="Backend API for College Report Generation and Approval System",
    version="1.0.0"
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)
APP_DIR = Path(__file__).resolve().parent


# ============================================================
# DATABASE CONFIGURATION
# ============================================================

DATABASE_URL = "sqlite:///./report_system.db"

engine = create_engine(
    DATABASE_URL,
    connect_args={"check_same_thread": False}
)

SessionLocal = sessionmaker(
    autocommit=False,
    autoflush=False,
    bind=engine
)

Base = declarative_base()

# Opaque, server-side login sessions keep approval authorization out of
# client-controlled role fields. Sessions expire when the API process restarts.
active_sessions: dict[str, int] = {}


# ============================================================
# DATABASE MODELS
# ============================================================

# -------------------------
# USER TABLE
# -------------------------

class User(Base):

    __tablename__ = "users"

    id = Column(
        Integer,
        primary_key=True,
        index=True
    )

    username = Column(
        String,
        unique=True,
        index=True,
        nullable=False
    )

    email = Column(
        String,
        unique=True,
        index=True,
        nullable=False
    )

    password = Column(
        String,
        nullable=False
    )

    role = Column(
        String,
        default="student",
        nullable=False
    )

    name = Column(String, nullable=True)
    college_code = Column(String, nullable=True)
    authorization_code = Column(String, nullable=True)


# -------------------------
# REPORT TABLE
# -------------------------

class Report(Base):

    __tablename__ = "reports"

    id = Column(
        Integer,
        primary_key=True,
        index=True
    )

    title = Column(
        String,
        index=True,
        nullable=False
    )

    department = Column(
        String,
        index=True,
        nullable=False
    )

    report_type = Column(
        String,
        index=True,
        nullable=False
    )

    description = Column(
        Text,
        nullable=False
    )

    created_by = Column(
        Integer,
        nullable=False
    )

    status = Column(
        String,
        default="Pending",
        index=True
    )

    approver_comment = Column(
        Text,
        nullable=True
    )

    report_date = Column(String, nullable=True)
    submitted_by = Column(String, nullable=True)
    submitted_role = Column(String, nullable=True)
    attachment_id = Column(String, nullable=True)
    file_name = Column(String, nullable=True)
    file_type = Column(String, nullable=True)
    event_id = Column(Integer, nullable=True)
    event_outcome = Column(String, nullable=True)
    submitted_on = Column(String, nullable=True)
    submitted_email = Column(String, nullable=True)
    approved_by_email = Column(String, nullable=True)


class AcademicEvent(Base):
    __tablename__ = "academic_events"
    id = Column(Integer, primary_key=True, index=True)
    academic_year = Column(String, nullable=False, index=True)
    event_name = Column(String, nullable=False)
    title = Column(String, nullable=False, unique=True)
    event_date = Column(String, nullable=False)
    celebrated = Column(Integer, default=1, nullable=False)
    celebration_status = Column(String, default="Pending", nullable=False)
    not_celebrated_reason = Column(Text, nullable=True)
    created_by = Column(Integer, nullable=False)


class AcademicCalendarFile(Base):
    __tablename__ = "academic_calendar_files"
    id = Column(Integer, primary_key=True)
    file_name = Column(String, nullable=False)
    content_type = Column(String, nullable=False)
    file_data = Column(LargeBinary, nullable=False)
    uploaded_by = Column(Integer, nullable=False)
    uploaded_on = Column(String, nullable=False)


# Create database tables

Base.metadata.create_all(bind=engine)

# Add fields used by the current frontend when upgrading an existing database.
with engine.begin() as connection:
    existing_columns = {column["name"] for column in inspect(engine).get_columns("users")}
    for column_name in ("name", "college_code", "authorization_code"):
        if column_name not in existing_columns:
            connection.execute(text(f"ALTER TABLE users ADD COLUMN {column_name} VARCHAR"))
    existing_report_columns = {column["name"] for column in inspect(engine).get_columns("reports")}
    for column_name in ("report_date", "submitted_by", "submitted_role", "attachment_id", "file_name", "file_type"):
        if column_name not in existing_report_columns:
            connection.execute(text(f"ALTER TABLE reports ADD COLUMN {column_name} VARCHAR"))
    if "event_id" not in existing_report_columns:
        connection.execute(text("ALTER TABLE reports ADD COLUMN event_id INTEGER"))
    if "submitted_on" not in existing_report_columns:
        connection.execute(text("ALTER TABLE reports ADD COLUMN submitted_on VARCHAR"))
    for column_name in ("submitted_email", "approved_by_email"):
        if column_name not in existing_report_columns:
            connection.execute(text(f"ALTER TABLE reports ADD COLUMN {column_name} VARCHAR"))
    if "event_outcome" not in existing_report_columns:
        connection.execute(text("ALTER TABLE reports ADD COLUMN event_outcome VARCHAR"))
    existing_event_columns = {column["name"] for column in inspect(engine).get_columns("academic_events")}
    if "celebration_status" not in existing_event_columns:
        connection.execute(text(
            "ALTER TABLE academic_events ADD COLUMN celebration_status VARCHAR NOT NULL DEFAULT 'Pending'"
        ))
        connection.execute(text(
            "UPDATE academic_events SET celebration_status = "
            "CASE WHEN celebrated = 1 THEN 'Celebrated' ELSE 'Not celebrated' END"
        ))


# ============================================================
# DATABASE CONNECTION
# ============================================================

def get_db():

    db = SessionLocal()

    try:

        yield db

    finally:

        db.close()


# ============================================================
# PYDANTIC SCHEMAS
# ============================================================


# -------------------------
# USER CREATE
# -------------------------

class UserCreate(BaseModel):

    username: str

    email: str

    password: str

    role: str
    name: str
    college_code: str
    authorization_code: str


# -------------------------
# USER RESPONSE
# -------------------------

class UserResponse(BaseModel):

    id: int

    username: str

    email: str

    role: str
    name: Optional[str] = None
    college_code: Optional[str] = None

    class Config:
        orm_mode = True


# -------------------------
# LOGIN
# -------------------------

class LoginRequest(BaseModel):

    username: str

    password: str
    role: str
    college_code: str
    authorization_code: str


# -------------------------
# REPORT CREATE
# -------------------------

class ReportCreate(BaseModel):

    title: str

    department: str

    report_type: str

    description: str

    created_by: int
    report_date: Optional[str] = None
    submitted_by: Optional[str] = None
    submitted_role: Optional[str] = None
    attachment_id: Optional[str] = None
    file_name: Optional[str] = None
    file_type: Optional[str] = None
    event_id: Optional[int] = None
    event_outcome: Optional[Literal["Organized", "Not organized"]] = None
    submitted_on: Optional[str] = None
    submitted_email: Optional[str] = None
    approved_by_email: Optional[str] = None


# -------------------------
# REPORT UPDATE
# -------------------------

class ReportUpdate(BaseModel):

    title: str

    department: str

    report_type: str

    description: str


# -------------------------
# REPORT RESPONSE
# -------------------------

class ReportResponse(BaseModel):

    id: int

    title: str

    department: str

    report_type: str

    description: str

    created_by: int

    status: str

    approver_comment: Optional[str] = None
    report_date: Optional[str] = None
    submitted_by: Optional[str] = None
    submitted_role: Optional[str] = None
    attachment_id: Optional[str] = None
    file_name: Optional[str] = None
    file_type: Optional[str] = None
    event_id: Optional[int] = None
    event_outcome: Optional[str] = None
    submitted_on: Optional[str] = None
    submitted_email: Optional[str] = None
    approved_by_email: Optional[str] = None

    class Config:
        orm_mode = True


# -------------------------
# APPROVAL REQUEST
# -------------------------

class ApprovalRequest(BaseModel):

    status: str

    approver_comment: Optional[str] = None


class AcademicEventCreate(BaseModel):
    academic_year: str
    event_name: str
    event_date: str


class AcademicEventCelebrationUpdate(BaseModel):
    celebrated: bool
    not_celebrated_reason: Optional[str] = None


class AcademicEventResponse(BaseModel):
    id: int
    academic_year: str
    event_name: str
    title: str
    event_date: str
    celebrated: bool
    celebration_status: str
    not_celebrated_reason: Optional[str] = None
    class Config:
        orm_mode = True


def get_current_user(
    authorization: Optional[str] = Header(default=None),
    db: Session = Depends(get_db)
):
    scheme, _, token = (authorization or "").partition(" ")
    user_id = active_sessions.get(token) if scheme.lower() == "bearer" else None
    user = db.query(User).filter(User.id == user_id).first() if user_id else None
    if user is None:
        raise HTTPException(status_code=401, detail="Please log in again.")
    return user


@app.get("/academic-events/", response_model=list[AcademicEventResponse])
def get_academic_events(db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    return db.query(AcademicEvent).order_by(AcademicEvent.event_date.desc()).all()


@app.get("/academic-calendar/file")
def get_academic_calendar_file(db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    document = db.query(AcademicCalendarFile).order_by(AcademicCalendarFile.id.desc()).first()
    if document is None:
        return {"available": False}
    return {"available": True, "file_name": document.file_name, "uploaded_on": document.uploaded_on}


def _normalize_calendar_header(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(value or "").strip().lower()).strip()


def _parse_calendar_date(value: object) -> Optional[date]:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raw = str(value or "").strip()
    for date_format in (
        "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%m/%d/%Y", "%d.%m.%Y",
        "%d-%m-%y", "%d/%m/%y", "%d %B %Y", "%d %b %Y",
        "%B %d %Y", "%B %d, %Y", "%b %d %Y", "%b %d, %Y",
    ):
        try:
            return datetime.strptime(raw, date_format).date()
        except ValueError:
            continue
    return None


def _calendar_header_indexes(rows: list[tuple]) -> Optional[tuple[int, int, Optional[int], int]]:
    date_headers = {"date", "event date", "date of event", "scheduled date", "day", "event date time"}
    name_headers = {"event", "event name", "name", "programme", "program", "activity",
                    "event activity", "title", "occasion", "festival", "name of event",
                    "name of programme", "name of program"}
    year_headers = {"academic year", "year", "session"}
    for header_index, row in enumerate(rows[:20]):
        headers = [_normalize_calendar_header(value) for value in row]
        date_column = next((i for i, value in enumerate(headers) if value in date_headers), None)
        name_column = next((i for i, value in enumerate(headers) if value in name_headers), None)
        year_column = next((i for i, value in enumerate(headers) if value in year_headers), None)
        if date_column is not None and name_column is not None and date_column != name_column:
            return header_index, name_column, year_column, date_column
    return None


def _extract_calendar_table_events(rows: list[tuple], datemode: Optional[int] = None) -> list[dict]:
    columns = _calendar_header_indexes(rows)
    if columns is None:
        return []
    header_index, name_column, year_column, date_column = columns
    extracted = []
    for row in rows[header_index + 1:]:
        if max(name_column, date_column, year_column or 0) >= len(row):
            continue
        name = str(row[name_column] or "").strip()
        date_value = row[date_column]
        parsed_date = _parse_calendar_date(date_value)
        if parsed_date is None and datemode is not None and isinstance(date_value, (int, float)):
            try:
                import xlrd
                from xlrd.biffh import XLRDError
                parsed_date = xlrd.xldate.xldate_as_datetime(date_value, datemode).date()
            except (ValueError, OverflowError, XLRDError):
                parsed_date = None
        if not name or parsed_date is None:
            continue
        year = str(row[year_column]).strip() if year_column is not None and row[year_column] else str(parsed_date.year)
        extracted.append({"event_name": name, "event_date": parsed_date.isoformat(), "academic_year": year})
    return extracted


def _extract_calendar_pdf_events(file_data: bytes) -> list[dict]:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        text_content = "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(file_data)).pages)
    except PdfReadError as error:
        raise ValueError("The PDF could not be read. Upload a valid, text-based PDF.") from error
    date_pattern = re.compile(
        r"\b(?:\d{4}-\d{1,2}-\d{1,2}|\d{1,2}[./-]\d{1,2}[./-]\d{2,4}|"
        r"\d{1,2}\s+[A-Za-z]{3,9}\s+\d{4}|[A-Za-z]{3,9}\s+\d{1,2},?\s+\d{4})\b"
    )
    lines = [re.sub(r"\s+", " ", line).strip(" \t|-\u2022") for line in text_content.splitlines()]
    extracted = []
    for index, line in enumerate(lines):
        match = date_pattern.search(line)
        if not match:
            continue
        parsed_date = _parse_calendar_date(match.group(0))
        if parsed_date is None:
            continue
        event_name = (line[:match.start()] + " " + line[match.end():]).strip(" \t|:-\u2022")
        if not event_name and index:
            previous_line = lines[index - 1]
            if previous_line and not date_pattern.search(previous_line):
                event_name = previous_line
        event_name = re.sub(r"^\s*(?:\d+[\.\)]?\s+)+", "", event_name).strip()
        if not event_name or _normalize_calendar_header(event_name) in {
            "academic calendar", "event", "event name", "date", "event date"
        }:
            continue
        extracted.append({
            "event_name": event_name,
            "event_date": parsed_date.isoformat(),
            "academic_year": str(parsed_date.year),
        })
    return extracted


def extract_academic_calendar_events(file_name: str, file_data: bytes) -> list[dict]:
    extension = Path(file_name).suffix.lower()
    if extension == ".csv":
        try:
            content = file_data.decode("utf-8-sig")
        except UnicodeDecodeError:
            content = file_data.decode("cp1252")
        try:
            dialect = csv.Sniffer().sniff(content[:8192], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        rows = [tuple(row) for row in csv.reader(StringIO(content), dialect)]
        return _extract_calendar_table_events(rows)
    if extension == ".xlsx":
        from openpyxl import load_workbook
        from openpyxl.utils.exceptions import InvalidFileException

        try:
            workbook = load_workbook(BytesIO(file_data), read_only=True, data_only=True)
        except (InvalidFileException, BadZipFile, OSError) as error:
            raise ValueError("The Excel workbook could not be read.") from error
        try:
            return [
                event
                for sheet in workbook.worksheets
                for event in _extract_calendar_table_events(list(sheet.iter_rows(values_only=True)))
            ]
        finally:
            workbook.close()
    if extension == ".xls":
        import xlrd
        from xlrd.biffh import XLRDError

        try:
            workbook = xlrd.open_workbook(file_contents=file_data)
        except XLRDError as error:
            raise ValueError("The Excel workbook could not be read.") from error
        extracted = []
        for sheet in workbook.sheets():
            rows = [tuple(sheet.row_values(row_index)) for row_index in range(sheet.nrows)]
            extracted.extend(_extract_calendar_table_events(rows, workbook.datemode))
        return extracted
    if extension == ".pdf":
        return _extract_calendar_pdf_events(file_data)
    raise HTTPException(status_code=400, detail="Unsupported academic calendar format.")


@app.put("/academic-calendar/file")
async def upload_academic_calendar_file(request: Request, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    if current_user.role != "Admin":
        raise HTTPException(status_code=403, detail="Only an Admin can upload or replace the academic calendar.")
    max_size = 25 * 1024 * 1024
    chunks = []
    size = 0
    async for chunk in request.stream():
        size += len(chunk)
        if size > max_size:
            raise HTTPException(status_code=413, detail="Academic calendar files must be 25 MB or smaller.")
        chunks.append(chunk)
    if not size:
        raise HTTPException(status_code=400, detail="Choose an academic calendar file to upload.")
    raw_name = unquote(request.headers.get("x-file-name", "academic-calendar"))
    safe_name = raw_name.replace("\\", "/").split("/")[-1].strip()
    extension = Path(safe_name).suffix.lower()
    allowed_extensions = {".pdf", ".xlsx", ".xls", ".csv"}
    if extension not in allowed_extensions:
        raise HTTPException(status_code=400, detail="Upload a PDF, Excel, or CSV calendar file.")
    file_data = b"".join(chunks)
    try:
        extracted_events = extract_academic_calendar_events(safe_name, file_data)
    except ModuleNotFoundError as error:
        raise HTTPException(
            status_code=503,
            detail="Calendar analysis is not installed on this server. Run 'pip install -r requirements.txt' and restart the app.",
        ) from error
    except (ValueError, csv.Error, UnicodeError, OSError, BadZipFile) as error:
        raise HTTPException(status_code=422, detail=f"Could not read the calendar file: {error}")
    if not extracted_events:
        raise HTTPException(
            status_code=422,
            detail="No events with recognizable names and dates were found. Use table columns named Event (or Event Name) and Date; scanned PDFs are not supported.",
        )

    content_type = request.headers.get("content-type", "application/octet-stream").split(";")[0]
    document = AcademicCalendarFile(file_name=safe_name, content_type=content_type,
        file_data=file_data, uploaded_by=current_user.id, uploaded_on=date.today().isoformat())
    events_added = 0
    events_updated = 0
    imported_titles = set()
    for event_data in extracted_events:
        year = event_data["academic_year"].strip()
        name = event_data["event_name"].strip()
        title = f"{name} ({year})"
        if title in imported_titles:
            continue
        imported_titles.add(title)
        existing = db.query(AcademicEvent).filter(AcademicEvent.title == title).first()
        if existing is None:
            db.add(AcademicEvent(
                academic_year=year,
                event_name=name,
                title=title,
                event_date=event_data["event_date"],
                celebrated=0,
                celebration_status="Pending",
                not_celebrated_reason=None,
                created_by=current_user.id,
            ))
            events_added += 1
        elif existing.event_date != event_data["event_date"]:
            existing.event_date = event_data["event_date"]
            existing.celebrated = 0
            existing.celebration_status = "Pending"
            existing.not_celebrated_reason = None
            events_updated += 1

    db.query(AcademicCalendarFile).delete()
    db.add(document)
    db.commit()
    return {
        "message": "Academic calendar uploaded and analyzed.",
        "file_name": safe_name,
        "uploaded_on": document.uploaded_on,
        "events_added": events_added,
        "events_updated": events_updated,
    }


@app.get("/academic-calendar/file/download")
def download_academic_calendar_file(db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    document = db.query(AcademicCalendarFile).order_by(AcademicCalendarFile.id.desc()).first()
    if document is None:
        raise HTTPException(status_code=404, detail="No academic calendar file has been uploaded.")
    return Response(content=document.file_data, media_type=document.content_type,
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(document.file_name)}"})


@app.post("/academic-events/", response_model=AcademicEventResponse)
def create_academic_event(event: AcademicEventCreate, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    if current_user.role != "Admin":
        raise HTTPException(status_code=403, detail="Only an Admin can manage the academic calendar.")
    try:
        event_date = date.fromisoformat(event.event_date)
    except ValueError:
        raise HTTPException(status_code=400, detail="Enter a valid event date.")
    year = event.academic_year.strip()
    name = event.event_name.strip()
    if not year or not name:
        raise HTTPException(status_code=400, detail="Academic year and event name are required.")
    title = f"{name} ({year})"
    if db.query(AcademicEvent).filter(AcademicEvent.title == title).first():
        raise HTTPException(status_code=409, detail="This event already exists for that academic year.")
    record = AcademicEvent(academic_year=year, event_name=name, title=title,
        event_date=event_date.isoformat(), celebrated=0, celebration_status="Pending",
        not_celebrated_reason=None,
        created_by=current_user.id)
    db.add(record)
    db.commit()
    db.refresh(record)
    return record


@app.put("/academic-events/{event_id}", response_model=AcademicEventResponse)
def update_academic_event(event_id: int, event: AcademicEventCreate, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    if current_user.role != "Admin":
        raise HTTPException(status_code=403, detail="Only an Admin can manage the academic calendar.")
    record = db.query(AcademicEvent).filter(AcademicEvent.id == event_id).first()
    if record is None:
        raise HTTPException(status_code=404, detail="Academic calendar event not found.")
    year = event.academic_year.strip()
    name = event.event_name.strip()
    if not year or not name:
        raise HTTPException(status_code=400, detail="Academic year and event name are required.")
    try:
        parsed_date = date.fromisoformat(event.event_date)
    except ValueError:
        raise HTTPException(status_code=400, detail="Enter a valid event date.")
    title = f"{name} ({year})"
    duplicate = db.query(AcademicEvent).filter(AcademicEvent.title == title, AcademicEvent.id != event_id).first()
    if duplicate:
        raise HTTPException(status_code=409, detail="This event already exists for that academic year.")
    record.academic_year = year
    record.event_name = name
    record.title = f"{record.event_name} ({record.academic_year})"
    updated_date = parsed_date.isoformat()
    if record.event_date != updated_date:
        record.celebrated = 0
        record.celebration_status = "Pending"
        record.not_celebrated_reason = None
    record.event_date = updated_date
    db.commit()
    db.refresh(record)
    return record


@app.put("/academic-events/{event_id}/celebration", response_model=AcademicEventResponse)
def update_academic_event_celebration(
    event_id: int,
    update: AcademicEventCelebrationUpdate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    if current_user.role != "Admin":
        raise HTTPException(status_code=403, detail="Only an Admin can update event celebration status.")
    record = db.query(AcademicEvent).filter(AcademicEvent.id == event_id).first()
    if record is None:
        raise HTTPException(status_code=404, detail="Academic calendar event not found.")
    if date.fromisoformat(record.event_date) >= date.today():
        raise HTTPException(status_code=400, detail="Event status can only be confirmed after the event day.")

    reason = (update.not_celebrated_reason or "").strip()
    if not update.celebrated and not reason:
        raise HTTPException(status_code=400, detail="Give a reason when an event was not celebrated.")
    record.celebrated = int(update.celebrated)
    record.celebration_status = "Celebrated" if update.celebrated else "Not celebrated"
    record.not_celebrated_reason = None if update.celebrated else reason
    db.commit()
    db.refresh(record)
    return record


# ============================================================
# HOME
# ============================================================

@app.get("/")
def home():
    return FileResponse(APP_DIR / "index.html")


@app.get("/index.html")
def login_page():
    return FileResponse(APP_DIR / "index.html")


@app.get("/home.html")
def dashboard_page():
    return FileResponse(APP_DIR / "home.html")


# ============================================================
# USER APIs
# ============================================================


# -------------------------
# CREATE USER
# -------------------------

@app.post(
    "/users/create/",
    response_model=UserResponse
)
def create_user(
    user: UserCreate,
    db: Session = Depends(get_db)
):

    # Check username

    existing_username = db.query(User).filter(
        User.username == user.username
    ).first()

    if existing_username:

        raise HTTPException(
            status_code=400,
            detail="Username already registered"
        )


    # Check email

    existing_email = db.query(User).filter(
        User.email == user.email
    ).first()

    if existing_email:

        raise HTTPException(
            status_code=400,
            detail="Email already registered"
        )


    # Check role

    allowed_roles = [
        "Admin",
        "Staff",
        "HOD",
        "Principal",
        "Authorized Person"
    ]

    if user.role not in allowed_roles:

        raise HTTPException(
            status_code=400,
            detail="Invalid role"
        )

    authorization_codes = {
        "Admin": "ADMIN2026",
        "Staff": "STAFF2026",
        "HOD": "HOD2026",
        "Principal": "PRINCIPAL2026",
        "Authorized Person": "AUTH2026"
    }
    if len(user.name.strip()) < 3 or len(user.college_code.strip()) < 4:
        raise HTTPException(status_code=400, detail="Please enter a valid name and college code.")
    if user.authorization_code.strip().upper() != authorization_codes[user.role]:
        raise HTTPException(status_code=400, detail="Invalid authorization code for selected role.")


    db_user = User(

        username=user.username,

        email=user.email,

        password=user.password,

        role=user.role,
        name=user.name.strip(),
        college_code=user.college_code.strip().upper(),
        authorization_code=user.authorization_code.strip().upper()

    )


    db.add(db_user)

    db.commit()

    db.refresh(db_user)

    return {
        "id": db_user.id,
        "username": db_user.username,
        "email": db_user.email,
        "role": db_user.role,
        "name": db_user.name,
        "college_code": db_user.college_code
    }


# -------------------------
# GET ALL USERS
# -------------------------

@app.get(
    "/users/",
    response_model=list[UserResponse]
)
def get_users(
    db: Session = Depends(get_db)
):

    users = db.query(User).all()

    return users


# -------------------------
# GET ONE USER
# -------------------------

@app.get(
    "/users/{user_id}",
    response_model=UserResponse
)
def get_user(
    user_id: int,
    db: Session = Depends(get_db)
):

    user = db.query(User).filter(
        User.id == user_id
    ).first()


    if user is None:

        raise HTTPException(
            status_code=404,
            detail="User not found"
        )


    return user


# ============================================================
# LOGIN
# ============================================================

@app.post("/users/login/")
def login_user(
    login: LoginRequest,
    db: Session = Depends(get_db)
):

    user = db.query(User).filter(User.email == login.username.lower()).first()


    if user is None:

        raise HTTPException(
            status_code=401,
            detail="Invalid username or password"
        )


    if (user.password != login.password
            or user.role != login.role
            or user.college_code != login.college_code.upper()
            or user.authorization_code != login.authorization_code.upper()):

        raise HTTPException(
            status_code=401,
            detail="Invalid username or password"
        )


    return {

        "message": "Login successful",

        "user_id": user.id,

        "username": user.username,

        "email": user.email,

        "role": user.role,
        "name": user.name or user.username,
        "college_code": user.college_code or "",
        "access_token": _create_session(user)

    }


def _create_session(user: User) -> str:
    token = secrets.token_urlsafe(32)
    active_sessions[token] = user.id
    return token


# ============================================================
# REPORT APIs
# ============================================================


# -------------------------
# CREATE REPORT
# -------------------------

@app.post(
    "/reports/",
    response_model=ReportResponse
)
def create_report(
    report: ReportCreate,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):

    # Check user

    user = db.query(User).filter(
        User.id == report.created_by
    ).first()


    if user is None:

        raise HTTPException(
            status_code=404,
            detail="User not found"
        )

    if current_user.id != report.created_by:
        raise HTTPException(status_code=403, detail="You can only submit reports for your own account.")
    if report.report_type == "Event Report" and not report.event_id:
        raise HTTPException(status_code=400, detail="Select the academic calendar event for this report.")
    if report.report_type == "Event Report" and not report.event_outcome:
        raise HTTPException(status_code=400, detail="Select whether the event was organized.")
    if report.report_type != "Event Report" and (report.event_id or report.event_outcome):
        raise HTTPException(status_code=400, detail="Event and event outcome are only valid for Event Reports.")
    if report.event_id:
        event = db.query(AcademicEvent).filter(AcademicEvent.id == report.event_id).first()
        if event is None:
            raise HTTPException(status_code=404, detail="Academic calendar event not found.")
        if report.event_outcome == "Organized" and event.celebration_status != "Celebrated":
            raise HTTPException(status_code=400, detail=f"{event.title} must be confirmed as celebrated before submitting an organized event report.")
        if report.event_outcome == "Not organized" and date.fromisoformat(event.event_date) >= date.today():
            raise HTTPException(status_code=400, detail="An event can only be reported as not organized after its scheduled day.")
        if report.event_outcome == "Not organized" and event.celebration_status == "Celebrated":
            raise HTTPException(status_code=400, detail=f"{event.title} has already been confirmed as celebrated.")
        if report.event_outcome == "Not organized" and not report.description.strip():
            raise HTTPException(status_code=400, detail="Explain why the event was not organized.")
        if report.report_date and report.report_date != event.event_date:
            raise HTTPException(status_code=400, detail="Report date must match the selected event date.")
        report.report_date = event.event_date


    # Create report

    db_report = Report(

        title=report.title,

        department=report.department,

        report_type=report.report_type,

        description=report.description,

        created_by=report.created_by,

        report_date=report.report_date,
        submitted_by=report.submitted_by,
        submitted_role=report.submitted_role,
        attachment_id=report.attachment_id,
        file_name=report.file_name,
        file_type=report.file_type,
        event_id=report.event_id,
        event_outcome=report.event_outcome,
        submitted_on=date.today().isoformat(),
        submitted_email=user.email,

        status="Pending"

    )


    db.add(db_report)

    db.commit()

    db.refresh(db_report)


    return db_report


# -------------------------
# GET ALL REPORTS
# -------------------------

@app.get(
    "/reports/",
    response_model=list[ReportResponse]
)
def get_reports(
    db: Session = Depends(get_db)
):

    reports = db.query(Report).order_by(
        Report.id.desc()
    ).all()


    return reports


# -------------------------
# GET ONE REPORT
# -------------------------

@app.get(
    "/reports/{report_id}",
    response_model=ReportResponse
)
def get_report(
    report_id: int,
    db: Session = Depends(get_db)
):

    report = db.query(Report).filter(
        Report.id == report_id
    ).first()


    if report is None:

        raise HTTPException(
            status_code=404,
            detail="Report not found"
        )


    return report


# ============================================================
# GET REPORTS CREATED BY USER
# ============================================================

@app.get(
    "/reports/user/{user_id}",
    response_model=list[ReportResponse]
)
def get_user_reports(
    user_id: int,
    db: Session = Depends(get_db)
):

    user = db.query(User).filter(
        User.id == user_id
    ).first()


    if user is None:

        raise HTTPException(
            status_code=404,
            detail="User not found"
        )


    reports = db.query(Report).filter(
        Report.created_by == user_id
    ).order_by(
        Report.id.desc()
    ).all()


    return reports


# ============================================================
# GET PENDING REPORTS
# ============================================================

@app.get(
    "/reports/status/pending",
    response_model=list[ReportResponse]
)
def get_pending_reports(
    db: Session = Depends(get_db)
):

    reports = db.query(Report).filter(
        Report.status == "Pending"
    ).order_by(
        Report.id.desc()
    ).all()


    return reports


# ============================================================
# GET APPROVED REPORTS
# ============================================================

@app.get(
    "/reports/status/approved",
    response_model=list[ReportResponse]
)
def get_approved_reports(
    db: Session = Depends(get_db)
):

    reports = db.query(Report).filter(
        Report.status == "Approved"
    ).order_by(
        Report.id.desc()
    ).all()


    return reports


# ============================================================
# GET REPORTS NEEDING QUERY REVIEW
# ============================================================

@app.get(
    "/reports/status/query-review",
    response_model=list[ReportResponse]
)
def get_query_review_reports(
    db: Session = Depends(get_db)
):

    reports = db.query(Report).filter(
        Report.status == "Query Review"
    ).order_by(
        Report.id.desc()
    ).all()


    return reports


# ============================================================
# UPDATE REPORT
# ============================================================

@app.put(
    "/reports/{report_id}",
    response_model=ReportResponse
)
def update_report(
    report_id: int,
    report_data: ReportUpdate,
    db: Session = Depends(get_db)
):

    report = db.query(Report).filter(
        Report.id == report_id
    ).first()


    if report is None:

        raise HTTPException(
            status_code=404,
            detail="Report not found"
        )


    # Update details

    report.title = report_data.title

    report.department = report_data.department

    report.report_type = report_data.report_type

    report.description = report_data.description


    db.commit()

    db.refresh(report)


    return report


# ============================================================
# APPROVE / REJECT REPORT
# ============================================================

@app.put(
    "/reports/{report_id}/approval",
    response_model=ReportResponse
)
def approve_or_reject_report(
    report_id: int,
    approval: ApprovalRequest,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):

    if current_user.role != "Authorized Person":
        raise HTTPException(status_code=403, detail="Only the Authorized Person can approve or reject files.")

    report = db.query(Report).filter(
        Report.id == report_id
    ).first()


    if report is None:

        raise HTTPException(
            status_code=404,
            detail="Report not found"
        )


    # Check status

    allowed_status = ["Approved", "Query Review"]


    if approval.status not in allowed_status:

        raise HTTPException(
            status_code=400,
            detail="Status must be Approved or Query Review."
        )

    if report.status != "Pending":
        raise HTTPException(status_code=409, detail="This report has already been reviewed.")


    # A query must include a question or requested correction.

    if (
        approval.status == "Query Review"
        and not approval.approver_comment
    ):

        raise HTTPException(
            status_code=400,
            detail="The query for the report is required."
        )


    report.status = approval.status

    report.approver_comment = approval.approver_comment
    report.approved_by_email = current_user.email if approval.status == "Approved" else None


    db.commit()

    db.refresh(report)


    return report


# ============================================================
# DELETE REPORT
# ============================================================

@app.delete(
    "/reports/{report_id}"
)
def delete_report(
    report_id: int,
    db: Session = Depends(get_db)
):

    report = db.query(Report).filter(
        Report.id == report_id
    ).first()


    if report is None:

        raise HTTPException(
            status_code=404,
            detail="Report not found"
        )


    db.delete(report)

    db.commit()


    return {

        "message": "Report deleted successfully",

        "report_id": report_id

    }


# ============================================================
# DASHBOARD STATISTICS
# ============================================================

@app.get("/dashboard/")
def dashboard(
    db: Session = Depends(get_db)
):

    total = db.query(Report).count()


    pending = db.query(Report).filter(
        Report.status == "Pending"
    ).count()


    approved = db.query(Report).filter(
        Report.status == "Approved"
    ).count()


    query_review = db.query(Report).filter(
        Report.status == "Query Review"
    ).count()


    users = db.query(User).count()


    return {

        "total_users": users,

        "total_reports": total,

        "pending_reports": pending,

        "approved_reports": approved,

        "query_review_reports": query_review

    }


# ============================================================
# RUN SERVER
# ============================================================

if __name__ == "__main__":

    import uvicorn

    uvicorn.run(
        "main:app",
        host="127.0.0.1",
        port=8000,
        reload=True
    )
